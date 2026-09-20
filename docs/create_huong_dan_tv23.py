# -*- coding: utf-8 -*-
"""Sinh 2 file Word hướng dẫn làm code nhận diện biển số."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = r"C:\Users\Admin\Downloads\OneDrive - Saigon University\Desktop\SMART_PARKING_ML\docs"
NAVY = RGBColor(0x0B, 0x3D, 0x6E)
TEAL = RGBColor(0x0D, 0x6E, 0x6E)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "0B3D6E"
ROW_ALT = "E8F1F8"


def set_run_font(run, size=13, bold=False, italic=False, color=DARK):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def para(doc, text, *, size=13, bold=False, italic=False, align="left", before=0, after=6, first=None, color=DARK):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first:
        p.paragraph_format.first_line_indent = Cm(first)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 13, bold=True, color=NAVY if level == 1 else TEAL)


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=13, bold=True)
        r2 = p.add_run(text)
        set_run_font(r2, size=13)
    else:
        r = p.add_run(text)
        set_run_font(r, size=13)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def cell_text(cell, text, *, bold=False, size=11, color=DARK, align="left", bg=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    if bg:
        shade(cell, bg)
    for para_ in cell.paragraphs:
        for run in para_.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.autofit = False
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, bold=True, size=11, color=WHITE, align="center", bg=HEADER_BG)
    for ri, row in enumerate(rows):
        bg = ROW_ALT if ri % 2 else "FFFFFF"
        for ci, val in enumerate(row):
            cell_text(t.rows[ri + 1].cells[ci], val, size=10, bg=bg)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return t


def setup(doc):
    s = doc.sections[0]
    s.page_width = Cm(21)
    s.page_height = Cm(29.7)
    s.left_margin = Cm(2)
    s.right_margin = Cm(1.8)
    s.top_margin = Cm(1.6)
    s.bottom_margin = Cm(1.6)
    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run("SMART PARKING ML  •  HƯỚNG DẪN LÀM CODE NHẬN DIỆN BIỂN SỐ")
    set_run_font(r, size=9, italic=True, color=NAVY)


def cover(doc, who, role):
    para(doc, "ĐỒ ÁN SMART PARKING MACHINE LEARNING", size=14, bold=True, align="center", after=2)
    para(doc, "HƯỚNG DẪN LÀM CODE — NHẬN DIỆN BIỂN SỐ XE", size=18, bold=True, align="center", after=4, color=NAVY)
    para(doc, who, size=16, bold=True, align="center", after=2, color=TEAL)
    para(doc, role, size=13, italic=True, align="center", after=10)
    para(
        doc,
        "Tài liệu này chỉ dành cho một thành viên. Làm đúng file được giao, commit Git thường xuyên, "
        "họp nhóm trưởng Châu khi xong mỗi mốc. Không xóa code người khác.",
        align="justify",
        first=1.0,
    )


def make_tv2():
    doc = Document()
    setup(doc)
    cover(
        doc,
        "DÀNH CHO THÀNH VIÊN 2",
        "Vai trò: Dữ liệu – chuẩn hóa biển số – đánh giá độ chính xác OCR",
    )

    heading(doc, "1. Bạn phụ trách gì trong phần nhận diện biển số?", 1)
    para(
        doc,
        "Phần nhận diện biển số gồm hai nửa: (A) đọc chữ từ ảnh, (B) đưa chữ đó vào hệ thống gửi xe. "
        "Thành viên 2 chịu trách nhiệm nửa A về mặt dữ liệu và chất lượng: biển số Việt Nam được chuẩn hóa thế nào, "
        "ảnh mẫu, nhãn đúng, độ chính xác OCR, các ca OCR nhầm (O/0, I/1). "
        "Bạn không làm giao diện Flask và không viết logic phạt/vé ra vào — đó là Thành viên 3.",
        align="justify",
        first=1.0,
    )

    heading(doc, "2. File bạn được phép sửa (và file không đụng)", 1)
    heading(doc, "2.1. File chính — bạn làm code ở đây", 2)
    table(
        doc,
        ["File", "Việc bạn phải làm"],
        [
            (
                "src/vision/plate_ocr.py",
                "Hoàn thiện normalize_plate, is_valid_vn_plate, plates_match, _assemble_from_items. "
                "Bổ sung quy tắc biển VN (ô tô 1 dòng, xe máy 2 dòng). Không xóa class PlateRecognizer.",
            ),
            (
                "src/vision/preprocess.py",
                "Cải thiện load_image (xoay EXIF), resize, CLAHE, khử nhiễu để OCR đọc rõ hơn trên ảnh tối/lóa.",
            ),
            (
                "src/vision/generate_plate_images.py",
                "Sinh thêm ảnh mẫu: nhiều font, nghiêng, mờ, 2 dòng. Hàm embed_plate_in_scene: biển dán vào ảnh xe giả lập.",
            ),
            (
                "src/vision/vehicle_type.py",
                "Hàm infer_vehicle_type_from_plate: 59B1… → Motorbike, 51A… → Car. Viết test cho các mẫu thật.",
            ),
            (
                "src/data/generate_data.py",
                "Cột license_plate phải khớp loại xe. Một sinh viên nên giữ một biển xuyên suốt (không random mỗi dòng).",
            ),
            (
                "src/data/clean_data.py",
                "Khi làm sạch CSV: normalize_plate, loại biển rỗng/sai format, thống nhất vehicle_type.",
            ),
            (
                "notebooks/11_plate_recognition.ipynb",
                "Notebook thí nghiệm: đo exact-match, character accuracy, confusion O/0. Xuất bảng số liệu cho báo cáo.",
            ),
            (
                "tests/test_plate_ocr.py",
                "Thêm test: biển có dấu chấm/gạch, 2 dòng, nhầm OCR, khớp/không khớp.",
            ),
            (
                "tests/test_data.py",
                "Test generate + clean có cột license_plate và biển sau clean không còn dấu cách.",
            ),
            (
                "data/external/plates/",
                "Thu thập/chụp ảnh thật (không biển người lạ nếu chưa che). Đặt tên file = biển chuẩn, ví dụ 59B112345.jpg.",
            ),
        ],
        [5.2, 12.0],
    )

    heading(doc, "2.2. File không sửa (tránh xung đột với TV3 và nhóm trưởng)", 2)
    bullet(doc, "app/templates/*.html, app/static/js/dashboard.js — giao diện thuộc TV3.")
    bullet(doc, "src/parking/penalty.py, logic checkout/phạt — TV3.")
    bullet(doc, "src/models/classification.py, regression.py, demand_prediction.py — mô hình thời gian gửi xe, không phải OCR.")
    bullet(doc, "Nếu cần API mới: viết đề xuất rồi nhờ Châu merge vào app/app.py, đừng tự xóa route có sẵn.")

    heading(doc, "3. Kiến thức cần nắm trước khi gõ code", 1)
    bullet(doc, "59-B1 123.45 (xe máy, thường 2 hàng) → chuẩn hóa 59B112345.", bold_prefix="Biển xe máy: ")
    bullet(doc, "51A-123.45 (ô tô, 1 hàng) → chuẩn hóa 51A12345.", bold_prefix="Biển ô tô: ")
    bullet(doc, "Bỏ hết dấu chấm, gạch, khoảng trắng; viết hoa; chỉ giữ A–Z và 0–9.", bold_prefix="Chuẩn hóa: ")
    bullet(doc, "Hai ký tự đầu là mã tỉnh (số). OCR hay đọc O thành 0 và ngược lại — bạn phải có bảng sửa có kiểm soát, không sửa chữ series (B trong 59B1) thành 8.", bold_prefix="Sửa OCR: ")
    bullet(doc, "Hai biển khớp khi normalize(a) == normalize(b).", bold_prefix="So khớp: ")
    bullet(doc, "Cắt vùng biển rồi mới OCR, không OCR cả bức ảnh mờ nếu đã có bbox.", bold_prefix="Pipeline: ")

    heading(doc, "4. Việc làm theo từng tuần (checklist code)", 1)
    heading(doc, "Tuần 1 — Chuẩn hóa và dữ liệu giả", 2)
    para(doc, "Mục tiêu: mọi biển trong CSV đều cùng một format, test xanh.", italic=True)
    bullet(doc, "Mở src/vision/plate_ocr.py, đọc normalize_plate. Viết thêm 10 ví dụ vào tests/test_plate_ocr.py.")
    bullet(doc, "Ví dụ test bắt buộc: '59-B1 123.45' == '59B112345'; '51A-123.45' == '51A12345'; '5O B112345' ra mã tỉnh 50 hoặc 59 tùy rule bạn chốt với Châu.")
    bullet(doc, "Hàm is_valid_vn_plate: reject 'ABC', '123', biển quá ngắn. Accept đủ 7–11 ký tự dạng tỉnh + series + số.")
    bullet(doc, "Hàm _assemble_from_items: OCR 2 dòng ['59-B1', '123.45'] phải ghép thành 59B112345.")
    bullet(doc, "src/data/generate_data.py: viết _random_license_plate. Nên tạo dict student_id → (plate, vehicle) rồi dùng lại, đừng random biển mỗi record.")
    bullet(doc, "Chạy: python -m unittest tests.test_plate_ocr tests.test_data -v")
    bullet(doc, "Sản phẩm tuần 1: test pass + đoạn mô tả 1 trang 'Quy tắc biển số VN dùng trong đồ án' gửi Châu đưa vào báo cáo.")

    heading(doc, "Tuần 2 — Ảnh mẫu và tiền xử lý", 2)
    para(doc, "Mục tiêu: có bộ ảnh để đo OCR, preprocess giúp đọc rõ hơn.", italic=True)
    bullet(doc, "Chạy python src/vision/generate_plate_images.py — kiểm tra thư mục data/external/plates/ có PNG biển và scene_59B112345.jpg.")
    bullet(doc, "Viết thêm 5–10 biến thể: biển nghiêng 10 độ, giảm sáng, thêm nhiễu Gaussian (dùng PIL/OpenCV). Lưu vào data/external/plates/hard/.")
    bullet(doc, "Tạo file data/external/plates/labels.csv với cột: filename, plate_true, vehicle_type, note.")
    bullet(doc, "Trong preprocess.py: thử kernel bilateral khác; ghi nhận xét ảnh nào rõ hơn vào notebook 11.")
    bullet(doc, "Không commit ảnh người lạ rõ mặt nếu chưa che. Ưu tiên ảnh tự chụp biển xe nhà / ảnh synthetic.")

    heading(doc, "Tuần 3 — Đánh giá OCR (số liệu cho báo cáo)", 2)
    para(doc, "Mục tiêu: bảng metric, không chỉ 'chạy được demo'.", italic=True)
    bullet(doc, "Trong notebook 11: vòng lặp từng dòng labels.csv, gọi recognize_plate(đường_dẫn_ảnh).")
    bullet(doc, "Tính Exact Match = (số biển đọc đúng hoàn toàn) / (tổng ảnh).")
    bullet(doc, "Tính Character Accuracy = trung bình số ký tự đúng vị trí / độ dài biển thật.")
    bullet(doc, "Liệt kê 10 ca sai: ảnh, biển thật, biển đọc, engine (rapidocr).")
    bullet(doc, "So sánh OCR ảnh crop sẵn với OCR ảnh scene (biển nằm trong ảnh lớn).")
    bullet(doc, "Ghi kết quả ra outputs/plate_ocr_metrics.csv để TV3 và Châu lấy vào slide.")
    bullet(doc, "Cài engine: pip install rapidocr-onnxruntime pillow opencv-python-headless")

    heading(doc, "Tuần 4 — Cứng hóa code và bàn giao", 2)
    bullet(doc, "Nếu OCR đọc dính chữ rác (VN, HONDA): lọc candidate bằng is_valid_vn_plate trước khi chọn best.")
    bullet(doc, "Không hard-code một biển mẫu trong hàm nhận diện. Mọi biển phải ra từ ảnh hoặc từ nhãn.")
    bullet(doc, "README ngắn 15 dòng trong data/external/plates/README.md: cách đặt tên file, cách chạy generate, cách đánh giá.")
    bullet(doc, "Pull request / báo Châu: danh sách file đổi + số Exact Match.")
    bullet(doc, "Diễn tập 3 phút: giải thích normalize và 1 confusion matrix OCR.")

    heading(doc, "5. Hướng dẫn đọc và sửa từng hàm (chi tiết code)", 1)
    heading(doc, "5.1. normalize_plate(raw)", 2)
    para(
        doc,
        "Input có thể là None, '59-b1 123.45', '59.B1.12345'. Output chỉ còn A–Z0–9, uppercase. "
        "Sửa OCR ở vị trí số (index 0–1 và từ index 4 trở đi). Không đổi ký tự series (thường index 2 là chữ). "
        "Sau khi sửa, tự gọi hàm trên 20 chuỗi và in ra, đừng đoán.",
        align="justify",
        first=1.0,
    )
    heading(doc, "5.2. plates_match(a, b)", 2)
    para(
        doc,
        "Dùng khi xe ra. TV3 sẽ gọi hàm này. Bạn phải đảm bảo '59-B1 123.45' khớp '59B112345' và không khớp biển khác. "
        "Chuỗi rỗng không được trả True.",
        align="justify",
        first=1.0,
    )
    heading(doc, "5.3. _assemble_from_items(items)", 2)
    para(
        doc,
        "items là list (text, confidence) từ OCR. Xe máy 2 dòng cần nối. Ưu tiên chuỗi đã valid. "
        "Nếu nối xong valid còn từng dòng thì không valid — chọn bản nối.",
        align="justify",
        first=1.0,
    )
    heading(doc, "5.4. generate_data — license_plate", 2)
    para(
        doc,
        "Ô tô: 2 số tỉnh + 1 chữ + 5 số. Xe máy: 2 số tỉnh + chữ+số (B1) + 5 số. "
        "Xác suất tỉnh: 29, 30, 50, 51, 59, 92. Gắn biển vào dict trước vòng lặp record.",
        align="justify",
        first=1.0,
    )
    heading(doc, "5.5. clean_parking_data", 2)
    para(
        doc,
        "Map cột license_plate qua normalize_plate. Drop dòng biển rỗng. "
        "Không drop hết dữ liệu vì regex quá chặt — nếu drop > 5% phải nới is_valid hoặc sửa generate.",
        align="justify",
        first=1.0,
    )

    heading(doc, "6. Lệnh bạn phải chạy được (copy nguyên)", 1)
    para(doc, "Mở PowerShell, cd vào thư mục đồ án (folder SMART_PARKING_ML), rồi:", after=4)
    para(doc, "pip install -r requirements.txt", size=12, italic=True)
    para(doc, "python -m unittest tests.test_plate_ocr tests.test_data -v", size=12, italic=True)
    para(doc, "python src/vision/generate_plate_images.py", size=12, italic=True)
    para(
        doc,
        "python -c \"from src.vision.plate_ocr import recognize_plate; print(recognize_plate(r'data/external/plates/scene_59B112345.jpg'))\"",
        size=11,
        italic=True,
    )

    heading(doc, "7. Sản phẩm bàn giao (Châu sẽ kiểm)", 1)
    table(
        doc,
        ["Sản phẩm", "Đạt khi"],
        [
            ("Bộ test unittest", "python -m unittest tests.test_plate_ocr tests.test_data -v → OK"),
            ("labels.csv + ảnh", "Ít nhất 20 ảnh (synthetic được), mỗi dòng có biển thật"),
            ("Bảng metric", "outputs/plate_ocr_metrics.csv có exact_match, char_acc"),
            ("Notebook 11", "Chạy từ trên xuống không lỗi (trừ khi thiếu ảnh)"),
            ("Đoạn báo cáo", "1–2 trang: quy tắc biển + kết quả OCR + hạn chế"),
        ],
        [5.5, 11.7],
    )

    heading(doc, "8. Khi kẹt thì hỏi Châu / TV3 thế này", 1)
    bullet(doc, "OCR engine none / không đọc được: đã pip install rapidocr-onnxruntime chưa? Có đang chạy đúng thư mục project không?")
    bullet(doc, "TV3 bảo biển trên web sai format: bạn soi normalize_plate, đừng bảo TV3 tự cắt chuỗi trên HTML.")
    bullet(doc, "Cần thêm API xuất metric: nhắn Châu, không tự xóa /api/scan-plate.")
    bullet(doc, "Conflict Git trên plate_ocr.py: giữ hàm recognize của nhóm, chỉ merge phần normalize/test của bạn.")

    heading(doc, "9. Cam kết", 1)
    para(
        doc,
        "Tôi (Thành viên 2) làm đúng phạm vi dữ liệu – chuẩn hóa – đánh giá OCR; không đụng UI/phạt; "
        "báo tiến độ mỗi tuần cho nhóm trưởng Châu.",
        align="justify",
        first=1.0,
        after=14,
    )
    para(doc, "Họ tên: ……………………     MSSV: ……………………     Ngày: …… / …… / 2026", align="center")
    para(doc, "Chữ ký: ……………………", align="center", before=20)

    path = BASE + r"\Huong_dan_TV2_Nhan_dien_bien_so.docx"
    doc.save(path)
    return path


def make_tv3():
    doc = Document()
    setup(doc)
    cover(
        doc,
        "DÀNH CHO THÀNH VIÊN 3",
        "Vai trò: Ảnh vào → tìm biển → giao diện cổng → lưu vé → đối chiếu lúc ra",
    )

    heading(doc, "1. Bạn phụ trách gì trong phần nhận diện biển số?", 1)
    para(
        doc,
        "Thành viên 3 biến kết quả OCR thành nghiệp vụ bãi xe. Bạn nhận 1 file ảnh có chứa biển số, "
        "gọi pipeline tìm vùng biển + OCR, hiện kết quả lên web, ghi biển vào CSDL, "
        "và lúc xe ra thì so biển/loại xe với lúc vào, tính phạt quá giờ. "
        "Bạn không viết thuật toán chuẩn hóa biển hay bảng metric OCR — đó là Thành viên 2. "
        "Bạn không train mô hình nhóm thời gian gửi xe — Châu tích hợp.",
        align="justify",
        first=1.0,
    )

    heading(doc, "2. File bạn được phép sửa", 1)
    heading(doc, "2.1. File chính", 2)
    table(
        doc,
        ["File", "Việc bạn phải làm"],
        [
            (
                "src/vision/plate_detect.py",
                "find_plate_regions: tìm hình chữ nhật giống biển trong ảnh xe. annotate_plate: vẽ khung xanh + chữ biển.",
            ),
            (
                "app/app.py",
                "Giữ và hoàn thiện /api/scan-plate (nhận 1 ảnh, nhận diện, lưu), /api/scans, /api/checkin, /api/checkout. "
                "Không xóa /api/predict.",
            ),
            (
                "src/parking/tickets.py",
                "Bảng plate_scans + tickets. Hàm save_plate_scan, list_scans, open_ticket, close_ticket, find_open_by_plate.",
            ),
            (
                "src/parking/penalty.py",
                "compute_penalty: phút thực tế vs phút ML dự đoán, ân hạn 15 phút, bậc tiền phạt.",
            ),
            (
                "app/templates/parking.html",
                "Trang '1 ảnh → nhận diện → lưu'. Preview ảnh, hiện crop, hiện biển lớn, bảng lịch sử scan.",
            ),
            (
                "app/templates/index.html",
                "Form vào/ra: chọn ảnh thì tự điền biển. Ra bãi: so khớp, hiện phạt, cảnh báo sai biển/sai loại xe.",
            ),
            (
                "app/templates/dashboard.html",
                "Danh sách xe đang trong bãi (status=open) lấy /api/tickets?status=open.",
            ),
            (
                "app/templates/history.html",
                "Lịch sử ra vào: biển vào, biển ra, khớp?, phạt.",
            ),
            (
                "app/static/js/dashboard.js",
                "Fetch API, render bảng, xử lý lỗi mạng.",
            ),
            (
                "app/static/css/style.css",
                "Nav, bảng, màu cảnh báo. Không phá layout Bootstrap trang chủ.",
            ),
        ],
        [5.2, 12.0],
    )

    heading(doc, "2.2. File không sửa", 2)
    bullet(doc, "src/vision/plate_ocr.py phần normalize_plate / is_valid_vn_plate / test metric — TV2.")
    bullet(doc, "src/data/generate_data.py, clean_data.py, notebooks đánh giá OCR — TV2.")
    bullet(doc, "src/models/*.py — mô hình thời gian gửi xe.")
    bullet(doc, "Được import hàm của TV2: normalize_plate, recognize_plate, plates_match, infer_vehicle_type. Không copy-paste rồi viết lại một bản khác.")

    heading(doc, "3. Luồng bạn phải code cho chạy end-to-end", 1)
    table(
        doc,
        ["Bước", "Người dùng", "Code của bạn phải làm"],
        [
            ("1", "Chọn 1 file ảnh có biển số", "parking.html: input type=file, preview."),
            ("2", "Bấm Nhận diện và lưu", "POST /api/scan-plate, field name bắt buộc là image."),
            ("3", "Hệ thống tìm biển trong ảnh", "plate_detect.find_plate_regions → crop."),
            ("4", "Đọc chữ", "Gọi recognize_plate (TV2/Châu). Hiện plate, crop, ảnh khoanh vùng."),
            ("5", "Ghi lại biển", "Lưu ảnh vào data/processed/scans/, INSERT plate_scans, append CSV."),
            ("6", "Xe vào bãi", "POST /api/checkin: biển + MSSV + gọi predict ML → open_ticket."),
            ("7", "Xe ra bãi", "POST /api/checkout: OCR biển ra, plates_match, vehicle_types_match, compute_penalty."),
            ("8", "Sai biển / sai loại xe", "status=alert, không cho 'ảo' là đã khớp. Hiện chữ đỏ trên UI."),
        ],
        [2.2, 5.5, 9.5],
    )

    heading(doc, "4. Việc làm theo từng tuần (checklist code)", 1)
    heading(doc, "Tuần 1 — Cổng 1 ảnh chạy được", 2)
    bullet(doc, "Chạy python app/app.py, mở http://127.0.0.1:5000/parking")
    bullet(doc, "Upload data/external/plates/scene_59B112345.jpg (ảnh có chứa biển, không phải chỉ tấm biển).")
    bullet(doc, "Kỳ vọng: biển 59B112345 hiện to, có ảnh crop, dòng 'Đã lưu bản ghi #…'.")
    bullet(doc, "Nếu engine=none: pip install rapidocr-onnxruntime rồi restart app.")
    bullet(doc, "Sửa parking.html nếu nút bị disable, hoặc field ảnh không tên image.")
    bullet(doc, "Kiểm tra file mới trong data/processed/scans/ và data/processed/plate_scans.csv")
    bullet(doc, "Sản phẩm: screenshot 3 tấm (form, kết quả đúng, bảng đã lưu).")

    heading(doc, "Tuần 2 — Tìm vùng biển trên ảnh thật", 2)
    bullet(doc, "Mở src/vision/plate_detect.py. Hiểu 3 nhánh: Canny+contour, mask trắng/vàng, black-hat.")
    bullet(doc, "Chụp 3 ảnh xe máy (biển rõ, biển hơi xa, biển lệch góc). Không đưa ảnh người lạ rõ mặt lên Git nếu chưa che.")
    bullet(doc, "Nếu không khoanh đúng: nới aspect ratio (xe máy ~1.2–2.2, ô tô ~2–5.5), tăng max_candidates.")
    bullet(doc, "Hàm annotate_plate phải vẽ được bbox lên ảnh trả về annotated_jpeg_b64.")
    bullet(doc, "Không nhúng model YOLO nếu chưa thống nhất Châu (nặng). Ưu tiên OpenCV cho demo đồ án.")
    bullet(doc, "Sản phẩm: 3 ảnh annotated lưu outputs/ (không commit file quá lớn).")

    heading(doc, "Tuần 3 — Vào bãi / ra bãi / phạt / sai biển", 2)
    bullet(doc, "Trang chủ index.html: chọn ảnh vào → tự gọi /api/scan-plate → ô biển được điền.")
    bullet(doc, "Check-in tạo vé status=open, hiện khu đề xuất (từ /api/predict, Châu gắn model).")
    bullet(doc, "Check-out đúng biển, đúng loại xe → allowed=true, status=closed.")
    bullet(doc, "Check-out cố tình nhập biển khác hoặc ô lookup_plate: hiện 'Biển số lúc ra không khớp lúc vào.'")
    bullet(doc, "Đổi loại xe Motorbike → Car lúc ra: cảnh báo loại xe.")
    bullet(doc, "Giờ ra muộn hơn duration_minutes + 15 phút: penalty_vnd > 0. Đọc penalty.py, đừng hard-code tiền trên HTML.")
    bullet(doc, "Dashboard liệt kê vé open; History hiện khớp/phạt.")
    bullet(doc, "Sản phẩm: kịch bản demo 5 ca (mục 7) chạy live không sửa code giữa chừng.")

    heading(doc, "Tuần 4 — Cứng UI, lỗi, demo hội đồng", 2)
    bullet(doc, "Thông báo lỗi tiếng Việt: thiếu ảnh, không đọc được biển, không tìm thấy vé mở, file > 8MB.")
    bullet(doc, "app.config MAX_CONTENT_LENGTH đã 8MB — hiện alert nếu user chọn ảnh quá nặng.")
    bullet(doc, "Nav: Vào/Ra, Cổng camera, Dashboard, History — trang nào cũng vào được.")
    bullet(doc, "Không để console.log thừa; không để JSON thô làm UI chính (parking.html phải có chữ biển to).")
    bullet(doc, "Viết 1 trang 'Kịch bản demo' (mục 7) in mang theo. Diễn tập 5 phút phần demo của bạn.")
    bullet(doc, "Nếu OneDrive khóa SQLite: báo Châu, tạm Pause sync.")

    heading(doc, "5. Hướng dẫn sửa từng chỗ trong code", 1)
    heading(doc, "5.1. POST /api/scan-plate (app/app.py)", 2)
    para(
        doc,
        "Nhận request.files['image']. Nếu thiếu file → 400. Gọi recognize_plate(bytes). "
        "Lưu ảnh gốc + crop JPEG vào data/processed/scans/ với timestamp. "
        "Gọi ticket_store.save_plate_scan(...). Trả JSON: plate, valid, confidence, engine, crop_jpeg_b64, "
        "annotated_jpeg_b64, saved, image_path. Không trả traceback HTML.",
        align="justify",
        first=1.0,
    )
    heading(doc, "5.2. plate_detect.find_plate_regions", 2)
    para(
        doc,
        "Input ảnh BGR. Output list dict {bbox: [x,y,w,h], crop, score}, điểm cao trước. "
        "Nếu không thấy gì, vẫn trả 1 vùng = cả ảnh để OCR thử. Đừng raise exception trên ảnh lạ.",
        align="justify",
        first=1.0,
    )
    heading(doc, "5.3. tickets.save_plate_scan", 2)
    para(
        doc,
        "INSERT plate_scans. Đồng thời ghi CSV (header nếu file chưa có). "
        "list_scans() ORDER BY id DESC. Dashboard/parking gọi /api/scans.",
        align="justify",
        first=1.0,
    )
    heading(doc, "5.4. /api/checkin và /api/checkout", 2)
    para(
        doc,
        "Checkin: biển từ OCR hoặc form plate; infer_vehicle_type; run_prediction; open_ticket. "
        "Checkout: tìm vé open theo ticket_id, lookup_plate, hoặc biển lúc ra. "
        "plates_match + vehicle_types_match + compute_penalty. "
        "status alert nếu sai biển/loại; closed nếu khớp. allowed = plate_ok and vehicle_ok.",
        align="justify",
        first=1.0,
    )
    heading(doc, "5.5. parking.html (bắt buộc làm đẹp đủ demo)", 2)
    bullet(doc, "input file + img preview khi chọn file.")
    bullet(doc, "Nút đổi text 'Đang nhận diện...' và disabled khi đang gửi.")
    bullet(doc, "Hiện plate to, valid, confidence, đường dẫn file đã lưu.")
    bullet(doc, "Hai ảnh: annotated và crop (data:image/jpeg;base64,...).")
    bullet(doc, "Bảng lịch sử từ GET /api/scans khi load trang và sau mỗi lần scan.")

    heading(doc, "6. Lệnh và URL", 1)
    para(doc, "cd vào folder SMART_PARKING_ML rồi:", after=4)
    para(doc, "pip install -r requirements.txt", size=12, italic=True)
    para(doc, "python app/app.py", size=12, italic=True)
    para(doc, "Trình duyệt: http://127.0.0.1:5000/parking", size=12, italic=True)
    para(doc, "Trang vào/ra: http://127.0.0.1:5000/", size=12, italic=True)
    para(doc, "Dashboard: /dashboard    Lịch sử: /history", size=12, italic=True)
    para(
        doc,
        "Ảnh thử: data/external/plates/scene_59B112345.jpg và 59B112345.png",
        size=12,
        italic=True,
    )

    heading(doc, "7. Năm ca demo bạn phải chạy được (in mang theo)", 1)
    table(
        doc,
        ["Ca", "Thao tác", "Kết quả bắt buộc thấy"],
        [
            ("A. Đọc ảnh có biển", "Upload scene_59B112345.jpg tại /parking", "Biển 59B112345, đã lưu #id"),
            ("B. Vào bãi", "Trang chủ, ảnh vào hoặc nhập biển, bấm vào bãi", "Vé + khu đề xuất + biển trên vé"),
            ("C. Ra đúng", "Cùng biển, cùng loại xe", "Cho phép ra, khớp biển/loại"),
            ("D. Ra sai biển", "lookup_plate = biển vào, plate ra khác", "Cảnh báo không khớp, status alert"),
            ("E. Quá giờ", "exit_time muộn hơn dự đoán + 15 phút", "penalty_vnd > 0, dòng quá giờ"),
        ],
        [3.5, 6.5, 7.2],
    )

    heading(doc, "8. Lỗi thường gặp và cách xử lý (bạn chịu trách nhiệm UI/API)", 1)
    bullet(doc, "Trang trắng / 500: xem terminal Flask. Thường do import src khi chạy sai thư mục — phải chạy python app/app.py từ root project.")
    bullet(doc, "database is locked: OneDrive đang sync parking.db — Pause sync.")
    bullet(doc, "Không đọc được biển: ảnh tối/mờ/che. Chụp lại; nhờ TV2 cải thiện preprocess, bạn đừng tự viết normalize khác.")
    bullet(doc, "Check-out 404: chưa check-in hoặc biển normalize khác. In ra plate lúc vào và lúc ra.")
    bullet(doc, "Model ML missing: Châu lo file .pkl. UI vẫn phải hiện vé; app.py đã có nhánh model_warning.")

    heading(doc, "9. Sản phẩm bàn giao (Châu kiểm)", 1)
    table(
        doc,
        ["Sản phẩm", "Đạt khi"],
        [
            ("Cổng /parking", "1 ảnh vào → biển hiện → file + DB có bản ghi"),
            ("Vào/Ra", "5 ca mục 7 chạy trên máy khác (không chỉ máy bạn)"),
            ("Dashboard + History", "Số liệu khớp vé vừa tạo"),
            ("Screenshot / clip 1–2 phút", "Để chiếu nếu demo live hỏng"),
            ("Đoạn báo cáo", "1–2 trang: luồng vào-ra, đối chiếu biển, phạt, hạn chế camera"),
        ],
        [5.5, 11.7],
    )

    heading(doc, "10. Cam kết", 1)
    para(
        doc,
        "Tôi (Thành viên 3) làm giao diện cổng, API vào/ra, lưu biển, đối chiếu và phạt; "
        "không tự ý đổi quy tắc chuẩn hóa biển của Thành viên 2; báo tiến độ mỗi tuần cho nhóm trưởng Châu.",
        align="justify",
        first=1.0,
        after=14,
    )
    para(doc, "Họ tên: ……………………     MSSV: ……………………     Ngày: …… / …… / 2026", align="center")
    para(doc, "Chữ ký: ……………………", align="center", before=20)

    path = BASE + r"\Huong_dan_TV3_Nhan_dien_bien_so.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(make_tv2())
    print(make_tv3())
